#!/usr/bin/env python3
"""Build the PP-XAI Assignment 1 PowerPoint deck (student ID only)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
FIGS = ROOT / "assets" / "figures"
DIAG = ROOT / "assets" / "diagrams"
OUT = ROOT / "202440724_PPXAI_Presentation.pptx"

# Theme
TEAL = RGBColor(0x0D, 0x73, 0x77)
SLATE = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF7, 0xF9, 0xFA)
ACCENT = RGBColor(0x1A, 0x7A, 0x4C)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _set_run(run, size=18, bold=False, color=SLATE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                 color=SLATE, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def _add_bullets(slide, left, top, width, height, items, size=16, color=SLATE, spacing=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = f"•  {item}"
        _set_run(run, size=size, color=color)
    return box


def _add_bar(slide, color=TEAL, height=Inches(0.08)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), SLIDE_W, height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _add_footer(slide, page, total=14):
    _add_textbox(
        slide, Inches(0.5), Inches(7.1), Inches(10), Inches(0.3),
        "PP-XAI  ·  Student ID 202440724  ·  MSc Applied AI  ·  University of Hull",
        size=11, color=MUTED,
    )
    _add_textbox(
        slide, Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3),
        f"{page} / {total}",
        size=11, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def _notes(slide, text):
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = text


def _blank(prs):
    blank = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(blank)


def _picture(slide, path, left, top, width=None, height=None):
    kwargs = {"left": left, "top": top}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), **kwargs)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ----- 1 Title -----
    s = _blank(prs)
    _add_bar(s, TEAL, Inches(0.12))
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()

    _add_textbox(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.4),
                 "MSc Applied Artificial Intelligence  ·  Dissertation Project",
                 size=16, color=TEAL, bold=True)
    _add_textbox(s, Inches(0.8), Inches(2.1), Inches(11.5), Inches(1.6),
                 "Privacy-Preserving Explainable AI for\nBuilding Energy Performance Prediction",
                 size=34, bold=True, color=SLATE)
    _add_textbox(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(0.6),
                 "A Federated Learning Approach with SHAP and LIME Integration (PP-XAI)",
                 size=20, color=MUTED)
    _add_textbox(s, Inches(0.8), Inches(5.0), Inches(11.5), Inches(0.8),
                 "Student ID: 202440724\nSupervisor: Mona  ·  University of Hull",
                 size=18, color=SLATE)
    _add_footer(s, 1)
    _notes(s,
           "Introduce yourself by student ID only if required for anonymity. "
           "State the project in one sentence: privacy-preserving explainable AI "
           "for high-rise / large-building energy prediction using federated learning "
           "plus SHAP and LIME. (~45–60s)")

    # ----- 2 Hook / Problem -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "The Problem", size=30, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.1), Inches(6.2), Inches(4.5), [
        "UK net-zero depends on better building energy insight",
        "Energy Performance Certificates (EPCs) are rich open data for ML",
        "High-rise and large commercial assets raise privacy and governance barriers",
        "Pooling microdata across councils / landlords is often restricted",
        "Centralised ML assumes data can be shared — often it cannot",
    ], size=19, spacing=14)
    _picture(s, DIAG / "privacy_accuracy_xai_triangle.png",
             Inches(7.0), Inches(1.0), width=Inches(5.8))
    _add_footer(s, 2)
    _notes(s,
           "Hook with the tension: valuable open EPC data vs privacy walls for "
           "sensitive assets. Point to the design triangle — privacy, accuracy, "
           "explainability must be co-designed. (~60–75s)")

    # ----- 3 Aim & RQs -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Aim, Research Questions & Objectives", size=28, bold=True, color=SLATE)
    _add_textbox(s, Inches(0.6), Inches(1.0), Inches(12), Inches(0.7),
                 "Aim: Develop and evaluate a privacy-preserving explainable AI framework "
                 "(FedAvg + SHAP/LIME) for high-rise building energy prediction on UK EPC data.",
                 size=18, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.85), Inches(12), Inches(4.2), [
        "RQ1 — Can federated learning approach centralised accuracy while keeping data local?",
        "RQ2 — How stable are SHAP/LIME explanations across federated vs centralised models?",
        "RQ3 — Which features dominate, and are they consistent across cities?",
        "RQ4 — How can federated XAI support transparency (EU AI Act context)?",
        "Objectives: centralised baselines → FedAvg (3 cities) → SHAP/LIME → stats → Streamlit prototype",
    ], size=18, spacing=12)
    _add_footer(s, 3)
    _notes(s,
           "State the aim once, then walk RQ1–RQ4 quickly. Emphasise that the "
           "project jointly evaluates accuracy AND explanation stability — not "
           "accuracy alone. (~60s)")

    # ----- 4 Options analysis -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.45),
                 "Options Analysis — What Was Considered & Why", size=28, bold=True, color=SLATE)
    _picture(s, DIAG / "options_analysis.png",
             Inches(0.5), Inches(0.8), width=Inches(12.3))
    _add_footer(s, 4)
    _notes(s,
           "REQUIRED section. Walk the decision matrix: compared centralised and FL; "
           "FedAvg primary (FedProx essentially tied); MLP for FL because trees cannot "
           "be FedAvg-averaged; SHAP+LIME; DP/secure aggregation deferred as future work. "
           "(~75–90s)")

    # ----- 5 Data -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Data & High-Rise Definition", size=28, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.0), Inches(6.5), Inches(4.8), [
        "Sources: MHCLG domestic & non-domestic EPCs (OGL v3.0)",
        "Scanned: ~13.8M domestic + ~841K non-domestic rows",
        "Filtered corpus: 5,663 records (London / Manchester / Birmingham)",
        "Domestic: flats with storey count ≥ 5 (honest mid/high-rise proxy)",
        "Non-domestic: floor area ≥ 5,000 m² (excl. schools/hospitals/etc.)",
        "Critical honesty: 99.4% commercial; only 36 residential flats",
        "Leakage controls: excluded CO₂ and asset/env score features",
    ], size=17, spacing=10)
    _picture(s, FIGS / "eda_city_counts.png",
             Inches(7.3), Inches(1.0), width=Inches(5.4))
    _picture(s, FIGS / "partition_typology_mix.png",
             Inches(7.3), Inches(4.0), width=Inches(5.4))
    _add_footer(s, 5)
    _notes(s,
           "Be explicit about the commercial dominance — examiners reward honesty. "
           "flat_storey_count ≥ 10 yields ~0 rows; ≥5 is a documented proxy. "
           "Mention leakage exclusion briefly. (~75s)")

    # ----- 6 Pipeline -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.45),
                 "PP-XAI Pipeline — What Was Developed", size=28, bold=True, color=SLATE)
    _picture(s, FIGS / "fig3_1_ppxai_pipeline.png",
             Inches(0.6), Inches(0.85), width=Inches(12.1))
    _add_footer(s, 6)
    _notes(s,
           "Walk left-to-right: ingest → filter → preprocess → centralised baselines "
           "AND federated MLP → SHAP/LIME → evaluation → Streamlit. One pipeline, "
           "two training regimes. (~60s)")

    # ----- 7 FL architecture -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.4),
                 "Federated Architecture & EnergyMLP", size=28, bold=True, color=SLATE)
    _picture(s, FIGS / "fig3_2_fedavg_architecture.png",
             Inches(0.4), Inches(0.75), width=Inches(7.2))
    _picture(s, FIGS / "fig3_3_mlp_architecture.png",
             Inches(7.7), Inches(1.2), width=Inches(5.2))
    _add_textbox(s, Inches(0.6), Inches(6.55), Inches(12), Inches(0.4),
                 "3 clients · 8 FedAvg rounds · 4 local epochs · sample-size weighted aggregation  "
                 "(McMahan et al., 2017; Beutel et al., 2020)",
                 size=15, color=MUTED)
    _add_footer(s, 7)
    _notes(s,
           "Only model updates move — not raw CSVs. Same MLP topology for central "
           "neural baseline and FL for a fair comparison. Trees remain centralised "
           "competitors. (~60s)")

    # ----- 8 Centralised results -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.4),
                 "Centralised Results — Gradient Boosting Leads", size=28, bold=True, color=SLATE)
    _picture(s, FIGS / "baseline_rmse.png",
             Inches(0.4), Inches(0.85), width=Inches(6.2))
    _picture(s, FIGS / "baseline_r2.png",
             Inches(6.8), Inches(0.85), width=Inches(6.0))
    _add_textbox(s, Inches(0.6), Inches(6.4), Inches(12), Inches(0.5),
                 "Best: Gradient Boosting — RMSE 113.04 · MAE 80.02 · R² 0.559  |  "
                 "Moderate accuracy without leakage features (credible, not inflated)",
                 size=16, color=SLATE)
    _add_footer(s, 8)
    _notes(s,
           "GB is best central model. Stress that excluding emissions features "
           "kept R² honest (~0.56 rather than spuriously >0.95). (~45–60s)")

    # ----- 9 Federated results -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.4),
                 "Federated Results (RQ1) — Near Parity with Central MLP", size=26, bold=True, color=SLATE)
    _picture(s, FIGS / "federated_convergence.png",
             Inches(0.4), Inches(0.8), width=Inches(7.5))
    _picture(s, DIAG / "results_headline.png",
             Inches(7.9), Inches(1.5), width=Inches(5.0))
    _add_textbox(s, Inches(0.6), Inches(6.45), Inches(12), Inches(0.4),
                 "FedAvg R8: RMSE 118.21 · R² 0.518  ≈  central MLP RMSE 117.04 · R² 0.528  "
                 "(H1: within ~1% RMSE of central MLP)",
                 size=16, color=SLATE)
    _add_footer(s, 9)
    _notes(s,
           "Show convergence from poor R1 to competitive R8. Headline: federated "
           "neural model tracks central neural accuracy without pooling data. (~60s)")

    # ----- 10 Critical stats -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Critical Analysis — Statistical vs Practical Gap", size=28, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.1), Inches(12), Inches(5.2), [
        "Wilcoxon (GB vs Fed absolute errors): p = 1.85×10⁻⁵ — statistically significant",
        "Paired t-test: p = 7.97×10⁻⁵",
        "Cohen’s d ≈ 0.059 — negligible effect size",
        "Bootstrap 95% CI RMSE overlap: GB 105.9–120.1 vs Fed 110.8–125.4",
        "Interpretation: gap is real but practically small under a privacy constraint",
        "Central GB still wins if pooling is allowed; FedAvg is the relevant benchmark when it is not",
    ], size=19, spacing=14)
    _add_footer(s, 10)
    _notes(s,
           "This is the critical-analysis slide for the 70–100% band. Do not hide "
           "significance — reframe with effect size and overlapping CIs. (~60–75s)")

    # ----- 11 XAI -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.25), Inches(12), Inches(0.4),
                 "Explainability (RQ2–RQ3) — Stable Across Regimes", size=26, bold=True, color=SLATE)
    _picture(s, FIGS / "shap_comparison.png",
             Inches(0.3), Inches(0.75), width=Inches(6.4))
    _picture(s, FIGS / "per_client_shap_comparison.png",
             Inches(6.8), Inches(0.75), width=Inches(6.1))
    _add_textbox(s, Inches(0.6), Inches(6.4), Inches(12), Inches(0.45),
                 "H2: Spearman ρ = 0.956 (central GB vs fed MLP)  ·  "
                 "H3: top-5 Jaccard avg 0.78 across cities  ·  "
                 "Drivers: storey_count, property type, natural gas  (Lundberg & Lee, 2017)",
                 size=15, color=SLATE)
    _add_footer(s, 11)
    _notes(s,
           "Privacy-preserving training did not arbitrarily reshuffle feature importance. "
           "Mention LIME complements SHAP for local stories (shown in demo). (~60s)")

    # ----- 12 Demo -----
    s = _blank(prs)
    _add_bar(s, ACCENT)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Live Demonstration — Streamlit Prototype", size=28, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.2), Inches(12), Inches(4.5), [
        "Interactive prediction of energy intensity (kWh/m²/year)",
        "Building inputs → preprocessor → best model → prediction",
        "Global SHAP importance table for transparency",
        "EU AI Act Article 13 transparency note embedded in the UI (H4)",
        "Demo command:  source scripts/env.sh  &&  streamlit run src/webapp/app.py",
    ], size=20, spacing=16)
    _add_textbox(s, Inches(0.6), Inches(5.8), Inches(12), Inches(0.8),
                 "Switch to the running app now — try 2–3 non-scripted buildings "
                 "(different cities / property types).",
                 size=18, bold=True, color=TEAL)
    _add_footer(s, 12)
    _notes(s,
           "PAUSE SLIDES. Demo live for ~2–3 minutes. Prefer interactive, "
           "non-scripted inputs. Show prediction + SHAP. Return to slides for "
           "discussion. (~120–180s)")

    # ----- 13 Discussion -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Discussion & Limitations", size=28, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.1), Inches(12), Inches(5.2), [
        "Trade-off: central GB wins on RMSE if pooling is allowed; FedAvg wins on data residency",
        "Explanation stability reduces fear that FL yields an ‘alien’ uninterpretable model",
        "Limitation: domestic high-rise labelling in bulk EPC is weak — commercial proxy dominates",
        "Limitation: in-process FedAvg simulator ≠ production secure FL (no DP / secure aggregation yet)",
        "Limitation: KernelSHAP uses subsamples; rankings can shift with background choice",
        "Strength: leakage ablation and schema honesty strengthen scientific credibility",
    ], size=18, spacing=12)
    _add_footer(s, 13)
    _notes(s,
           "Be candid. Examiners mark critical evaluation highly. Position limitations "
           "as methodological contributions, not excuses. (~75s)")

    # ----- 14 Conclusions -----
    s = _blank(prs)
    _add_bar(s)
    _add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                 "Conclusions, Contributions & Future Work", size=28, bold=True, color=SLATE)
    _add_bullets(s, Inches(0.6), Inches(1.0), Inches(12), Inches(4.5), [
        "Federated MLP approaches central neural accuracy (R² ≈ 0.52 vs 0.53) without pooling raw data",
        "SHAP rankings remain highly aligned across regimes (ρ ≈ 0.96); cross-city consistency holds",
        "Working Streamlit prototype demonstrates stakeholder-facing prediction + explanations",
        "Contributions: PP-XAI pipeline, FL↔XAI stability evidence, honest high-rise schema audit",
        "Future work: UPRN–OSM residential linkage, DP-SGD / secure aggregation, more clients",
    ], size=18, spacing=12)
    _add_textbox(s, Inches(0.6), Inches(5.7), Inches(12), Inches(0.6),
                 "Thank you — Questions welcome",
                 size=24, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    _add_textbox(s, Inches(0.6), Inches(6.35), Inches(12), Inches(0.4),
                 "Key refs: McMahan et al. (2017); Lundberg & Lee (2017); Ribeiro et al. (2016); "
                 "Beutel et al. (2020); UK EPC data under OGL v3.0",
                 size=13, color=MUTED, align=PP_ALIGN.CENTER)
    _add_footer(s, 14)
    _notes(s,
           "Close with the triangle message. Invite Q&A. Keep under 15 minutes total "
           "including demo. (~45–60s)")

    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
